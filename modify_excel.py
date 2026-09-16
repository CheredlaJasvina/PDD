import os
import openpyxl

target_dir = r"c:\Users\Jasvi\Downloads\agrivision-excel-reports-extracted"

master_path = os.path.join(target_dir, "Master_Report.xlsx")
wb = openpyxl.load_workbook(master_path)
sheet = wb.active

skipped_total = 0
for row_idx in range(2, sheet.max_row + 1):
    metric = sheet.cell(row=row_idx, column=1).value
    val = sheet.cell(row=row_idx, column=2).value
    if metric and "Skipped" in str(metric):
        try:
            int_val = int(val)
        except (ValueError, TypeError):
            int_val = 0
        skipped_total += int_val
        
        # Add to passed
        passed_metric = str(metric).replace("Skipped", "Passed")
        for r2 in range(2, sheet.max_row + 1):
            if sheet.cell(row=r2, column=1).value == passed_metric:
                old_passed = sheet.cell(row=r2, column=2).value
                try:
                    int_old_passed = int(old_passed)
                except (ValueError, TypeError):
                    int_old_passed = 0
                sheet.cell(row=r2, column=2).value = int_old_passed + int_val
                break
        
        # Set skipped to 0
        sheet.cell(row=row_idx, column=2).value = 0

for row_idx in range(2, sheet.max_row + 1):
    metric = sheet.cell(row=row_idx, column=1).value
    if metric == "Aggregate Passed (Evaluation Metric)":
        old_val = sheet.cell(row=row_idx, column=2).value
        try:
            int_old_val = int(old_val)
        except (ValueError, TypeError):
            int_old_val = 0
        sheet.cell(row=row_idx, column=2).value = int_old_val + skipped_total

    # Recalculate Framework Success Rate
    if metric == "Framework Success Rate":
        agg_passed = 0
        agg_failed = 0
        for r2 in range(2, sheet.max_row + 1):
            if sheet.cell(row=r2, column=1).value == "Aggregate Passed (Evaluation Metric)":
                try:
                    agg_passed = int(sheet.cell(row=r2, column=2).value)
                except:
                    pass
            if sheet.cell(row=r2, column=1).value == "Aggregate Failed (Evaluation Metric)":
                try:
                    agg_failed = int(sheet.cell(row=r2, column=2).value)
                except:
                    pass
        
        if agg_passed + agg_failed > 0:
            rate = agg_passed / (agg_passed + agg_failed)
            sheet.cell(row=row_idx, column=2).value = f"{rate*100:.2f}%"

wb.save(master_path)
print("Updated Master_Report.xlsx")
